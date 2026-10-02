"""Every leg of every route at or under the record, one sheet per route (docs/route_legs.xlsx), for block tests and
manual pruning. Same routes and evaluation as route_report.py, plus per-leg new states and flags.

Columns: depart / arrive with stop and state, service, ride hours, wait before, miss penalty (next departure from the
same stop that still reaches the same alighting stop: about +24 h is a once-daily anchor), new states, flags.

Leaf script: nothing imports it.   python3 route_sheets.py   (from scripts/)
"""
import pickle
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import time_search as ts
from route_report import R, miss_penalty, st
from common import ABBR

RED, GREEN, BLUE, AMBER = (PatternFill('solid', fgColor=c) for c in ('F8D7DA', 'D4EDDA', 'D6E4F0', 'FFF3CD'))


def legs_for(start, t0, order, end=None, fixed=True):
    S = ts.Search(R, start, t0, [], 1, end=end)
    ts._R, ts._START, ts._T0, ts._SELF, ts._BASE, ts._END = R, start, t0, S, None, end
    if not fixed:
        b = S.settle(order); start = b[1]; t0 = S.latest_start(start, b[2], order)
        S = ts.Search(R, start, t0, [], 1, end=end); ts._START, ts._T0, ts._SELF = start, t0, S
    f, legs = S.evaluate(order)
    assert S.all_states(legs), [R.state_names[s] for s in S.missing(legs)]
    return S, f, legs


def sheet(wb, name, S, f, legs):
    ws = wb.create_sheet(name)
    rides = [l for l in legs if l[0] == 'ride']; first = rides[0][3]
    ws.append([f'{name}: {(f - first) / 3600:.2f} h, first bus {R.fmt_time(first, rides[0][2])}, finish {R.fmt_time(f, rides[-1][4])}'])
    ws['A1'].font = Font(bold=True, size=12)
    ws.append(['Red: connection under 20 min. Green: wait over 2 h. Blue: hand-entered trip (data/manual). '
               'Amber: miss penalty of 20 h or more (a once-daily anchor).'])
    hdr = ['#', 'kind', 'depart', 'from', 'st', 'arrive', 'to', 'st', 'service', 'ride h', 'wait before h',
           'miss penalty h', 'new states', 'total states', 'flags']
    ws.append(hdr)
    for c in ws[3]: c.font = Font(bold=True)
    seen = set(S.done); s0 = R.state[legs[0][2] if legs[0][0] == 'ride' else legs[0][1]]
    if s0 >= 0: seen.add(int(s0))
    prev = None; n = 0
    for l in legs:
        n += 1
        if l[0] == 'walk':
            _, a, b, t0_, t1 = l
            new = set()
            if b >= R.first_virtual and R.state[b] >= 0: new = {int(R.state[b])} - seen
            seen |= new
            row = [n, 'walk', R.fmt_time(t0_, a), R.stop_name[a], R.abbr[a] or '', R.fmt_time(t1, b), R.stop_name[b], R.abbr[b] or '',
                   'walk / local', round((t1 - t0_) / 3600, 2), '', '', ' '.join(sorted(ABBR[R.state_names[s]] for s in new)), len(seen), '']
            ws.append(row); prev = t1; continue
        _, trip, bs, bt, al, at = l
        cov = S.covered([l]); new = cov - seen; seen |= cov
        ag, rt, hs = R.trip_info[(trip[0], trip[1])]
        w = None if prev is None else (bt - prev) / 3600
        mp = miss_penalty(bs, bt, al, at)
        manual = R.feeds[trip[0]].name == 'extra'
        flags = []
        if w is not None and w < 20 / 60: flags.append('tight')
        if w is not None and w > 2: flags.append('long wait')
        if manual: flags.append('manual')
        if mp is not None and mp >= 20: flags.append('daily anchor')
        ws.append([n, 'ride', R.fmt_time(bt, bs), R.stop_name[bs], R.abbr[bs] or '', R.fmt_time(at, al), R.stop_name[al], R.abbr[al] or '',
                   f'{ag}: {rt or hs}', round((at - bt) / 3600, 2), '' if w is None else round(w, 2),
                   'none in 72 h' if mp is None else round(mp, 1), ' '.join(sorted(ABBR[R.state_names[s]] for s in new)), len(seen), ', '.join(flags)])
        r = ws.max_row
        if 'tight' in flags: ws.cell(r, 11).fill = RED
        if 'long wait' in flags: ws.cell(r, 11).fill = GREEN
        if manual: ws.cell(r, 9).fill = BLUE
        if mp is not None and mp >= 20: ws.cell(r, 12).fill = AMBER
        prev = at
    for col, wdt in zip('ABCDEFGHIJKLMNO', (4, 6, 15, 36, 4, 15, 36, 4, 48, 7, 9, 10, 14, 7, 28)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = 'A4'


if __name__ == '__main__':
    wb = Workbook(); wb.remove(wb.active)
    w = pickle.load(open('../out/res_wheeling.pkl', 'rb'))
    sheet(wb, 'Wheeling 234.24', *legs_for(R.stop_name.index(w['start']), w['t0'], w['order']))
    c = pickle.load(open('../out/res_starttest.pkl', 'rb'))
    sheet(wb, 'Control 237.24', *legs_for(R.stop_name.index(c['start']), c['t0'], c['order']))
    lv = R.stops_matching('Primm Bus Stop (702 Shuttles)')[0]   # 702 Shuttles 10:50 (Oct 2 screenshot): 50 min better than the Las Vegas 10:00 Greyhound
    # pin Baton Rouge: Delta Bus Lines from Memphis (MS on the way), then the 23:55 Orlando bus straight to Tallahassee and Jacksonville (Nick, Oct 2). Ties the Lafayette backtrack, and is bookable
    rx = [st(a) if len(a) == 2 else a for a in 'NV,CA,AR,TN,@Baton Rouge Bus Station,LA,FL,VT,ME,KY,MI,ND,SD,OR,WY'.split(',')]
    sheet(wb, 'RX 234.75', *legs_for(lv, datetime(2026, 9, 23, 23, 0, tzinfo=ZoneInfo('America/Los_Angeles')).timestamp(), rx, end=st('WY'), fixed=False))
    mex = [x for x in c['order'] if x not in (st('VT'), st('ME'))] + ['@Hanover (Hanover Inn)', '@White River Junction Bus Station', '@Boston South Station', st('ME')]
    sheet(wb, 'MEX 237.75', *legs_for(R.stop_name.index(c['start']), c['t0'], mex, end=st('ME'), fixed=False))
    import re
    xtol = []
    for line in open('../docs/cloud_tol.txt', encoding='utf-8'):
        m = re.search(r'->  d\d+ \w+ \d\d:\d\d  (.*), (?:[A-Z]{2}|None)   \[', line)
        if m and re.match(r'^d\d+ ', line) and (not xtol or xtol[-1] != '@' + m.group(1)): xtol.append('@' + m.group(1))
    xtol.append(st('VT'))
    sheet(wb, 'X_TOL 237.91', *legs_for(R.stops_matching('Sheridan WY Taco Bell')[0], datetime(2026, 9, 24, tzinfo=ZoneInfo('America/Denver')).timestamp(), xtol, end=st('VT'), fixed=False))
    j = pickle.load(open('../out/shapes/join.pkl', 'rb'))
    sheet(wb, 'X 238.11', *legs_for(R.stop_name.index(j['start']), j['t0'], j['order'], end=st('MS')))
    wb.save('../docs/route_legs.xlsx')
    print('wrote ../docs/route_legs.xlsx:', ', '.join(ws.title for ws in wb.worksheets))
