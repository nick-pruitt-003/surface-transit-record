"""Export a timetabled result (from time_router/time_search) to the spreadsheet
format: one row per ride with day, local depart/arrive, board/alight, carrier.

usage: python3 time_sheet.py --result ../out/time_best.pkl --at 2026-11-02T07:00 --tz America/Denver --days 20 --out ../out/timed.xlsx
"""
import argparse, pickle
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from time_router import TimeRouter, load_feeds, STATE_TZ, WALK_MAX_KM
from common import hav, ABBR


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--result', required=True); ap.add_argument('--at', required=True); ap.add_argument('--tz', default='America/Denver')
    ap.add_argument('--days', type=int, default=20); ap.add_argument('--out', required=True); ap.add_argument('--title', default='Timetabled route')
    a = ap.parse_args()
    t0 = datetime.fromisoformat(a.at).replace(tzinfo=ZoneInfo(a.tz))
    R = TimeRouter(load_feeds(), t0.date(), a.days)
    res = pickle.load(open(a.result, 'rb')); legs = res['legs']
    def local(ts, stop):
        ab = R.abbr[stop] or 'TX'
        d = datetime.fromtimestamp(ts, ZoneInfo(STATE_TZ.get(ab, 'UTC')))
        return (d.date() - t0.date()).days + 1, d.strftime('%H:%M')
    wb = Workbook(); ws = wb.active; ws.title = 'Route'
    hdr = ['Leg', 'Day', 'Depart', 'Arrive', 'Arr day', 'Board at', 'Alight at', 'Carrier', 'Route', 'States on leg', 'New states', 'Wait before (h)', 'Price', 'Notes']
    ws.append(hdr)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
    seen = {R.abbr[[l for l in legs if l[0] == 'ride'][0][2]]}
    n = 0; prev_arr = t0.timestamp(); pending_note = ''
    for l in legs:
        if l[0] == 'walk':
            _, s1, s2, tw, _ = l; d = hav(R.stop_xy[s1], R.stop_xy[s2])
            if d > 0.05: pending_note = f"{'walk' if d <= WALK_MAX_KM else 'local transit'} {d:.1f} km {R.stop_name[s1]} -> {R.stop_name[s2]}; "
            continue
        _, trip, bs, bt, al, at = l
        fi, tid, _ = trip; f = R.feeds[fi]; t = f.trips[tid]
        ids = [R.idx[f.stops[s[1]][0]] for s in t['stops']]
        i, j = ids.index(bs), ids.index(al)
        sts = []
        for k in ids[min(i, j):max(i, j) + 1]:
            ab = R.abbr[k]
            if ab and (not sts or sts[-1] != ab): sts.append(ab)
        fresh = []
        for k in ids[min(i, j):max(i, j) + 1]:
            ab = R.abbr[k]
            if ab and ab not in seen: seen.add(ab); fresh.append(f"{ab} @ {R.stop_name[k]}")
        n += 1
        dday, dtime = local(bt, bs); aday, atime = local(at, al) if at else ('', '')
        ag, rt, hs = R.trip_info[(fi, tid)]
        ws.append([n, dday, dtime, atime, aday, f"{R.stop_name[bs]}, {R.abbr[bs] or ''}", f"{R.stop_name[al]}, {R.abbr[al] or ''}", ag, rt or hs, '>'.join(sts), ', '.join(fresh), round((bt - prev_arr) / 3600, 1), '', pending_note])
        if fresh:
            for c in ws[ws.max_row]: c.fill = PatternFill('solid', fgColor='E2EFDA')
        prev_arr = at or bt; pending_note = ''
    fin = res.get('finish'); 
    ws.append([]); ws.append(['Finish', *local(fin, [l for l in legs if l[0]=='ride'][-1][4]), '', '', f"elapsed {round((fin - t0.timestamp())/3600,1)} h = {round((fin - t0.timestamp())/86400,2)} days", '', '', '', '', f"{len(seen)} states"])
    for i, w in enumerate([5, 5, 8, 8, 7, 38, 38, 22, 40, 22, 40, 9, 8, 50], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws2 = wb.create_sheet('Notes')
    for line in [a.title, '', 'Times are local to each stop. Day 1 = start date. Green rows enter a new state.',
                 'Greyhound/FlixBus times from the current Flix GTFS; Amtrak from Amtrak GTFS; regional carriers from the Jan-2026 merged feed (weekday pattern only, verify).',
                 'Wait before = hours between the previous arrival and this departure.']:
        ws2.append([line])
    ws2.column_dimensions['A'].width = 120
    wb.save(a.out); print('saved', a.out, n, 'legs')


if __name__ == '__main__':
    main()
